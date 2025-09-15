import openpyxl

def read_excel_data(filename="TwilioSender.xlsx"):
    """
    Reads all data from an Excel workbook and prints it to the console.
    """
    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
    except FileNotFoundError:
        print(f"Error: The file '{filename}' was not found.")
        return

    print(f"--- Reading data from: {filename} ---\n")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"--- Sheet: {ws.title} ---")

        # Check if the sheet is empty
        if not ws.calculate_dimension() or (ws.calculate_dimension() == 'A1' and ws.cell(row=1, column=1).value is None):
            print("(Sheet is empty)")
            print("-" * (len(ws.title) + 14))
            print("\n")
            continue

        data_found = False
        max_cols = ws.max_column
        for row in ws.iter_rows():
            # Using a list comprehension for conciseness
            row_data = [str(cell.value) if cell.value is not None else "" for cell in row]

            # Only print rows that contain some data
            if any(row_data):
                print(" | ".join(row_data))
                data_found = True

        if not data_found:
            print("(Sheet has no data)")

        print("-" * (len(ws.title) + 14))
        print("\n")

if __name__ == "__main__":
    read_excel_data()
