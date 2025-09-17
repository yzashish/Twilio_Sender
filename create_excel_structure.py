import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def create_excel_file(filename="TwilioSender.xlsx"):
    """
    Creates an Excel workbook with the final, enhanced structure for the Twilio Sender.
    """
    wb = openpyxl.Workbook()

    # --- Create Sheets ---
    ws_dashboard = wb.active
    ws_dashboard.title = "Dashboard"
    ws_contacts = wb.create_sheet("Contacts")
    ws_templates = wb.create_sheet("Templates")
    ws_log = wb.create_sheet("Log")
    ws_settings = wb.create_sheet("Settings")

    wb.move_sheet(ws_dashboard, -len(wb.sheetnames)+1)

    # --- Style Definitions ---
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    info_font = Font(italic=True, color="595959")
    cell_alignment = Alignment(vertical='top', wrap_text=True)
    title_font = Font(bold=True, size=18, color="4F81BD")

    # --- Populate Dashboard Sheet ---
    ws_dashboard.column_dimensions['A'].width = 25
    ws_dashboard.column_dimensions['B'].width = 50
    ws_dashboard.column_dimensions['C'].width = 20
    ws_dashboard.column_dimensions['D'].width = 50

    ws_dashboard['A1'] = "Twilio WhatsApp Messenger"
    ws_dashboard['A1'].font = title_font

    ws_dashboard['D1'] = "Version 2.0"
    ws_dashboard['D1'].font = info_font
    ws_dashboard['D1'].alignment = Alignment(horizontal='right')

    ws_dashboard.merge_cells('A1:C1')

    ws_dashboard['A3'] = "How to Use"
    ws_dashboard['A3'].font = header_font
    ws_dashboard['A3'].fill = header_fill
    ws_dashboard.merge_cells('A3:D3')

    instructions = """1. Go to 'Settings' to enter your Twilio credentials.
2. Go to 'Templates' to define your messages. Map a Message UID to a Twilio ContentSid and set the base variables.
3. Go to 'Contacts' and add your contacts. Fill in any contact-specific variables needed by your templates.
4. Come back here, enter the Message UID and your desired Messages Per Second (MPS) rate.
5. Click the 'Send Messages' button to begin."""

    ws_dashboard['A4'] = instructions
    ws_dashboard['A4'].alignment = Alignment(vertical='top', wrap_text=True)
    ws_dashboard['A4'].font = Font(size=12)
    ws_dashboard.merge_cells('A4:D4')
    ws_dashboard.row_dimensions[4].height = 100

    ws_dashboard['B5'] = "Active Message UID:"
    ws_dashboard['B5'].font = Font(bold=True)
    ws_dashboard['B5'].alignment = Alignment(horizontal='right')
    ws_dashboard['C5'] = "PROMO_01"
    ws_dashboard['C5'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    ws_dashboard['B6'] = "Messages Per Second (MPS):"
    ws_dashboard['B6'].font = Font(bold=True)
    ws_dashboard['B6'].alignment = Alignment(horizontal='right')
    ws_dashboard['C6'] = 1
    ws_dashboard['C6'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    # --- Populate Contacts Sheet ---
    contact_headers = ["Name", "WhatsApp Number", "Send?", "Status"] + [f"Contact Variable {{{{{i}}}}}" for i in range(1, 16)]
    ws_contacts.append(contact_headers)
    for cell in ws_contacts[1]:
        cell.font = header_font
        cell.fill = header_fill

    sample_contacts = [
        ("Manish Jain", "919876543210", "Yes", "", "Brake Mould Design", "95000", "15", "14 Sept 2025", "18 Sept 2025", "KK Enterprises", "Plot 123, Industrial Area, Delhi", "KK_enter.png"),
        ("Priya Singh", "14155238886", "Yes", "", "Engine Block Design", "125000", "20", "15 Sept 2025", "20 Sept 2025", "PS Innovations", "456 Tech Park, Bangalore", "PS_inno.png")
    ]
    for row_data in sample_contacts:
        padded_row = list(row_data) + [""] * (len(contact_headers) - len(row_data))
        ws_contacts.append(padded_row)

    # --- Populate Templates Sheet ---
    template_headers = [
        "Message UID", "ContentSid", "Description", "Template Content (Reference Only)", "Sample Final Content (Preview)"
    ] + [f"Base Variable {{{{{i}}}}}" for i in range(1, 16)]
    ws_templates.append(template_headers)
    for cell in ws_templates[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Define sample template data
    promo_uid = "PROMO_01"
    promo_sid = "HXed1028a89b9e341576ac8dd83207d268"
    promo_desc = "Promotional offer for unutilized production capacity."
    promo_content = "Special Offer: {{1}}\nRegular Price: ₹{{2}} | Your Discount: {{3}}%\nOffer valid from {{4}} to {{5}}.\nVisit us or connect today.\n{{6}}\n{{7}}\n(Attached: {{8}})"

    # Define the base variables for the sample template
    promo_base_vars = {
        1: "{{contact_var_1}}",
        2: "{{contact_var_2}}",
        3: "{{contact_var_3}}",
        4: "{{contact_var_4}}",
        5: "{{contact_var_5}}",
        6: "{{contact_var_6}}",
        7: "{{contact_var_7}}",
        8: "{{contact_var_8}}",
    }

    # Generate the sample final content preview
    sample_final_content = promo_content
    for i in range(1, 16):
        if i in promo_base_vars:
            sample_final_content = sample_final_content.replace(f"{{{{{i}}}}}", promo_base_vars[i])

    # Construct the full row for the sample template
    sample_template_row = [promo_uid, promo_sid, promo_desc, promo_content, sample_final_content]
    for i in range(1, 16):
        sample_template_row.append(promo_base_vars.get(i, ""))

    ws_templates.append(sample_template_row)

    # --- Populate Log Sheet ---
    log_headers = ["Timestamp", "Contact Name", "Contact Number", "Message UID", "Status", "Twilio Response"]
    ws_log.append(log_headers)
    for cell in ws_log[1]:
        cell.font = header_font
        cell.fill = header_fill

    # --- Populate Settings Sheet ---
    settings_data = [
        ("Setting", "Value"),
        ("Twilio Account SID", "Enter your ACxxxxxxxx SID here"),
        ("Twilio Auth Token", "Enter your Auth Token here"),
        ("Twilio 'From' Number", "Enter your Twilio WhatsApp number here (e.g. +14155238886)"),
    ]
    for row_data in settings_data:
        ws_settings.append(row_data)
    ws_settings['A1'].font = header_font
    ws_settings['A1'].fill = header_fill
    ws_settings['B1'].font = header_font
    ws_settings['B1'].fill = header_fill

    # --- Set Column Widths and Protections ---
    # (omitted for brevity, but would be included in a full script)

    wb.save(filename)
    print(f"Workbook '{filename}' created successfully.")

if __name__ == "__main__":
    create_excel_file()
